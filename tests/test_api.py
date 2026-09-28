from __future__ import annotations

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import deps
from app.main import app
from app.services.gemini_client import GeminiRateLimitError
from app.services.glossary import GlossaryMatcher
from app.services.pipeline import TranslationPipeline
from tests.test_pipeline import FakeLLM, FakeTM


class RateLimitedLLM:
    async def generate_structured(self, **_: object) -> None:
        raise GeminiRateLimitError("Gemini rate limit reached. Please wait a moment and try again.")


@pytest.fixture
def client(matcher: GlossaryMatcher, monkeypatch: pytest.MonkeyPatch):  # type: ignore[no-untyped-def]
    llm = FakeLLM(by_source={"": "عدالت عالیہ نے ملزم کی ضمانت منظور کر لی۔"})
    monkeypatch.setattr(deps, "_pipeline", TranslationPipeline(llm=llm, glossary=matcher, tm=FakeTM()))
    monkeypatch.setattr(deps, "rate_limiter", deps.RateLimiter(3))
    with TestClient(app) as c:
        deps.glossary.set_terms(matcher.terms)
        yield c


def test_health(client: TestClient) -> None:
    r = client.get("/api/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "ok" and body["app"] == "Tashreeh AI" and body["supabase"] is False


def test_translate(client: TestClient) -> None:
    r = client.post("/api/translate", json={"text": "The High Court granted bail to the accused.", "direction": "en-ur"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["translation"].startswith("عدالت عالیہ")
    assert body["terms_missing"] == []
    assert len(body["glossary_matches"]) == 3
    assert body["saved"] is False and body["id"]


def test_empty_and_too_long_input(client: TestClient) -> None:
    r = client.post("/api/translate", json={"text": "   ", "direction": "en-ur"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "empty_input"
    r = client.post("/api/translate", json={"text": "a" * 6000, "direction": "en-ur"})
    assert r.status_code == 413 and r.json()["error"]["code"] == "input_too_long"


def test_bad_direction_and_mismatch(client: TestClient) -> None:
    r = client.post("/api/translate", json={"text": "hello", "direction": "en-fr"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_request"
    r = client.post("/api/translate", json={"text": "ملزم کی ضمانت", "direction": "en-ur"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "direction_mismatch"


def test_rate_limit(client: TestClient) -> None:
    payload = {"text": "bail", "direction": "en-ur", "session_id": "s1"}
    codes = [client.post("/api/translate", json=payload).status_code for _ in range(4)]
    assert codes == [200, 200, 200, 429]
    # Another session is unaffected.
    assert client.post("/api/translate", json={**payload, "session_id": "s2"}).status_code == 200


def test_gemini_rate_limit_is_reported(client: TestClient, matcher: GlossaryMatcher, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(deps, "_pipeline", TranslationPipeline(llm=RateLimitedLLM(), glossary=matcher, tm=FakeTM()))  # type: ignore[arg-type]
    r = client.post("/api/translate", json={"text": "bail", "direction": "en-ur"})
    assert r.status_code == 429 and r.json()["error"]["code"] == "gemini_rate_limited"


def test_glossary_list_and_filters(client: TestClient) -> None:
    body = client.get("/api/glossary").json()
    assert body["total"] == len(deps.glossary) and body["read_only"] is True
    assert body["category_counts"]["court"] == 4
    assert [t["term_en"] for t in client.get("/api/glossary", params={"q": "ضمانت"}).json()["items"]] == ["bail"]
    assert all(t["category"] == "land" for t in client.get("/api/glossary", params={"category": "land"}).json()["items"])


def test_glossary_write_needs_database(client: TestClient) -> None:
    r = client.post("/api/glossary", json={"term_en": "x", "term_ur": "ی"})
    assert r.status_code == 503 and r.json()["error"]["code"] == "read_only"


def test_history_and_feedback_need_database(client: TestClient) -> None:
    assert client.get("/api/history", params={"session_id": "s"}).status_code == 503
    r = client.post(
        "/api/feedback", json={"translation_id": "00000000-0000-0000-0000-000000000000", "rating": 1}
    )
    assert r.status_code == 503


def test_word_fix_alone_is_valid_feedback(client: TestClient) -> None:
    fix = {"term_id": 1, "source": "bail", "old_target": "ضمانت", "new_target": "بیل"}
    base = {"translation_id": "00000000-0000-0000-0000-000000000000"}
    # Gets past validation (and stops only because there is no database in tests).
    assert client.post("/api/feedback", json={**base, "term_corrections": [fix]}).status_code == 503
    unchanged = {**fix, "new_target": "ضمانت"}
    r = client.post("/api/feedback", json={**base, "term_corrections": [unchanged]})
    assert r.status_code == 422 and r.json()["error"]["code"] == "empty_feedback"


def test_review_and_glossary_edit_need_database(client: TestClient) -> None:
    assert client.get("/api/review").status_code == 503
    assert client.post("/api/review/1/approve", json={}).status_code == 503
    r = client.patch("/api/glossary/1", json={"term_ur_common": "بیل"})
    assert r.status_code == 503 and r.json()["error"]["code"] == "read_only"


def _docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    doc.add_heading("Order", level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "The accused"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_document_translation(client: TestClient) -> None:
    data = _docx_bytes(["The accused was granted bail.", "", "The witness appeared."])
    r = client.post(
        "/api/translate/document",
        files={"file": ("order.docx", data, "application/octet-stream")},
        data={"direction": "en-ur"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert [s["kind"] for s in body["segments"]] == ["heading", "paragraph", "paragraph", "table_cell"]
    assert body["docx_filename"] == "order.ur.docx" and body["docx_base64"]


def test_document_stream(client: TestClient) -> None:
    data = _docx_bytes(["The accused was granted bail."])
    r = client.post(
        "/api/translate/document?stream=true",
        files={"file": ("order.docx", data, "application/octet-stream")},
        data={"direction": "en-ur"},
    )
    lines = [l for l in r.text.splitlines() if l]
    import json

    events = [json.loads(l)["type"] for l in lines]
    assert events[0] == "start" and events[-1] == "done" and events.count("segment") == 3


def test_unsupported_and_scanned_files(client: TestClient) -> None:
    r = client.post("/api/translate/document", files={"file": ("a.txt", b"hello", "text/plain")}, data={"direction": "en-ur"})
    assert r.status_code == 422 and "Unsupported" in r.json()["error"]["message"]

    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    r = client.post(
        "/api/translate/document", files={"file": ("scan.pdf", buf.getvalue(), "application/pdf")}, data={"direction": "en-ur"}
    )
    assert r.status_code == 422 and "scanned" in r.json()["error"]["message"]


def test_export_docx(client: TestClient) -> None:
    r = client.post("/api/translate/export", json={"translation": "ملزم کی ضمانت", "direction": "en-ur"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    Document(io.BytesIO(r.content))  # parses
