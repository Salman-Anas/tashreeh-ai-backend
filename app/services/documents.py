"""DOCX / PDF text extraction and translated-DOCX rebuild (basic formatting only)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Literal

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.models import Direction

SegmentKind = Literal["paragraph", "heading", "table_cell"]

URDU_FONT = "Noto Nastaliq Urdu"
ENGLISH_FONT = "Calibri"

DISCLAIMER_EN = "Draft translation for assistance only. Verify with a qualified legal professional before official use."
DISCLAIMER_UR = "یہ ترجمہ صرف معاونت کے لیے ایک مسودہ ہے۔ سرکاری استعمال سے قبل کسی مستند قانونی ماہر سے تصدیق کر لیں۔"


class DocumentError(Exception):
    """User-facing extraction error (unsupported, scanned, empty, corrupt)."""


@dataclass
class Segment:
    kind: SegmentKind
    text: str
    heading_level: int = 0


@dataclass
class TableBlock:
    rows: list[list[int]]  # segment indexes (-1 = empty cell)


@dataclass
class ExtractedDocument:
    segments: list[Segment] = field(default_factory=list)
    # Layout: ("seg", index) for a paragraph/heading, ("table", TableBlock) for a table.
    layout: list[tuple[str, int | TableBlock]] = field(default_factory=list)

    @property
    def total_chars(self) -> int:
        return sum(len(s.text) for s in self.segments)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def _iter_blocks(doc: DocxDocument):  # type: ignore[no-untyped-def]
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _heading_level(p: Paragraph) -> int:
    name = (p.style.name if p.style is not None else "") or ""
    if name == "Title":
        return 1
    m = re.match(r"Heading (\d)", name)
    return int(m.group(1)) if m else 0


def extract_docx(data: bytes) -> ExtractedDocument:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError("This .docx file could not be opened. It may be corrupt or password-protected.") from exc

    out = ExtractedDocument()
    for block in _iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            level = _heading_level(block)
            out.segments.append(Segment("heading" if level else "paragraph", text, level))
            out.layout.append(("seg", len(out.segments) - 1))
        else:
            rows: list[list[int]] = []
            for row in block.rows:
                cells: list[int] = []
                seen: set[int] = set()
                for cell in row.cells:
                    # Merged cells repeat the same underlying element; keep one.
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    text = cell.text.strip()
                    if text:
                        out.segments.append(Segment("table_cell", text))
                        cells.append(len(out.segments) - 1)
                    else:
                        cells.append(-1)
                rows.append(cells)
            if rows:
                out.layout.append(("table", TableBlock(rows)))
    if not out.segments:
        raise DocumentError("No text found in this document.")
    return out


# ---------------------------------------------------------------------------
# PDF (text-based only)
# ---------------------------------------------------------------------------
_SENTENCE_END = re.compile(r"[.:;?!۔؟]\s*$")


def _pdf_paragraphs(page_text: str) -> list[str]:
    lines = [ln.strip() for ln in page_text.replace("\r", "\n").split("\n")]
    paras: list[str] = []
    buf: list[str] = []
    for ln in lines:
        if not ln:
            if buf:
                paras.append(" ".join(buf))
                buf = []
            continue
        buf.append(ln)
        if _SENTENCE_END.search(ln):
            paras.append(" ".join(buf))
            buf = []
    if buf:
        paras.append(" ".join(buf))
    return [p for p in paras if p]


def extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise DocumentError("This PDF is password-protected.") from exc
        pages = [page.extract_text() or "" for page in reader.pages]
    except DocumentError:
        raise
    except (PdfReadError, Exception) as exc:  # noqa: BLE001
        raise DocumentError("This PDF could not be read. It may be corrupt.") from exc

    total = sum(len(p.strip()) for p in pages)
    if not pages or total < 20 * len(pages):
        raise DocumentError(
            "This PDF appears to be scanned (image only) and has no selectable text. "
            "Scanned PDFs are not supported yet — please upload a .docx or a text-based PDF."
        )

    out = ExtractedDocument()
    for page_text in pages:
        for para in _pdf_paragraphs(page_text):
            out.segments.append(Segment("paragraph", para))
            out.layout.append(("seg", len(out.segments) - 1))
    return out


def extract(filename: str, data: bytes) -> ExtractedDocument:
    name = filename.lower()
    if name.endswith(".docx"):
        return extract_docx(data)
    if name.endswith(".pdf"):
        return extract_pdf(data)
    if name.endswith(".doc"):
        raise DocumentError("Old .doc files are not supported. Please save the file as .docx and try again.")
    raise DocumentError("Unsupported file type. Please upload a .docx or a text-based .pdf file.")


# ---------------------------------------------------------------------------
# DOCX rebuild
# ---------------------------------------------------------------------------
def _set_rtl(paragraph: Paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "right")
    ppr.append(jc)


def _style_run(run, rtl: bool, size: int) -> None:  # type: ignore[no-untyped-def]
    font = URDU_FONT if rtl else ENGLISH_FONT
    run.font.name = font
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), font)
    if rtl:
        rpr.append(OxmlElement("w:rtl"))
        cs = OxmlElement("w:szCs")
        cs.set(qn("w:val"), str(size * 2))
        rpr.append(cs)


def _add_text(container, text: str, rtl: bool, size: int, style: str | None = None) -> None:  # type: ignore[no-untyped-def]
    lines = text.split("\n")
    p = container.add_paragraph(style=style) if style else container.add_paragraph()
    if rtl:
        _set_rtl(p)
    for i, line in enumerate(lines):
        run = p.add_run(line)
        _style_run(run, rtl, size)
        if i < len(lines) - 1:
            run.add_break()


def build_docx(
    extracted: ExtractedDocument, translations: list[str], direction: Direction, title: str | None = None
) -> bytes:
    rtl = direction == "en-ur"
    doc = Document()
    if title:
        _add_text(doc, title, rtl, 16, style="Title")

    for kind, ref in extracted.layout:
        if kind == "seg":
            assert isinstance(ref, int)
            seg = extracted.segments[ref]
            text = translations[ref]
            if seg.kind == "heading":
                level = min(max(seg.heading_level, 1), 4)
                _add_text(doc, text, rtl, {1: 18, 2: 16, 3: 14}.get(level, 13), style=f"Heading {level}")
            else:
                _add_text(doc, text, rtl, 13 if rtl else 11)
        else:
            assert isinstance(ref, TableBlock)
            ncols = max(len(r) for r in ref.rows)
            table = doc.add_table(rows=len(ref.rows), cols=ncols)
            table.style = "Table Grid"
            if rtl:
                tblpr = table._tbl.tblPr
                bidi = OxmlElement("w:bidiVisual")
                tblpr.append(bidi)
            for r, row in enumerate(ref.rows):
                for c, seg_idx in enumerate(row):
                    cell = table.cell(r, c)
                    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
                    _add_text(cell, translations[seg_idx] if seg_idx >= 0 else "", rtl, 12 if rtl else 10)
            doc.add_paragraph()

    doc.add_paragraph()
    _add_text(doc, DISCLAIMER_EN, False, 8)
    _add_text(doc, DISCLAIMER_UR, True, 9)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_text_docx(translation: str, direction: Direction, source_text: str | None = None) -> bytes:
    """A simple DOCX of a single translation (optionally with the source appended)."""
    rtl = direction == "en-ur"
    doc = Document()
    for para in translation.split("\n\n"):
        _add_text(doc, para.strip(), rtl, 13 if rtl else 11)
    if source_text:
        doc.add_paragraph()
        _add_text(doc, "Source text", False, 10)
        for para in source_text.split("\n\n"):
            _add_text(doc, para.strip(), not rtl, 11 if rtl else 13)
    doc.add_paragraph()
    _add_text(doc, DISCLAIMER_EN, False, 8)
    _add_text(doc, DISCLAIMER_UR, True, 9)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
